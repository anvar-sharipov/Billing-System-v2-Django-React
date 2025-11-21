# core/views.py
import pandas as pd
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from django.core.files.storage import default_storage
import os
import uuid
from datetime import datetime
from icecream import ic

from rest_framework import viewsets
from .models import Etrap
from .serializers import EtrapSerializer

from rest_framework.decorators import api_view
from rest_framework.response import Response
from django.core.exceptions import ValidationError
import re
from .models import *
from django.core.paginator import Paginator, EmptyPage
from django.shortcuts import get_object_or_404
from .serializers import UserTableSerializer, UserDogoworSerializer
from django.utils.timezone import make_aware
from .views import get_group_list


from django.http import HttpResponse
from django.db.models import Q
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import io
from django.db.models import Subquery, OuterRef
from django.db.models import Count


current_year = datetime.now().year



@api_view(['GET'])
def checkActiveOrNot(request):
    try:
        # Получаем параметры из query string
        number = request.GET.get('number')
        etrap = request.GET.get('etrap')
        
        activate_at = request.GET.get('activate_at')
        deactivate_at = request.GET.get('deactivate_at')
        
        activate_at_datetime = None
        if activate_at:
            activate_at_datetime = datetime.fromisoformat(activate_at)
         
        deactivate_at_datetime = None
        if deactivate_at:
            deactivate_at_datetime = datetime.fromisoformat(deactivate_at)
            
        
        # ic(activate_at_datetime)
        # ic(deactivate_at_datetime)
        
        
        user_id = request.GET.get('id')

        # ic("Received parameters:", number, etrap)
        
        # Валидация обязательных параметров
        if not number:
            return Response({
                "success": False,
            }, status=400)
        
        if not etrap:
            return Response({
                "success": False, 
            }, status=400)
        
   
        
        if activate_at_datetime and not deactivate_at_datetime: 
            # Ищем пользователей с таким номером и этапом, у которых есть активные договоры
            query = UserTable.objects.filter(
                number=number,
                etrap_id=etrap,
                dogowors__activate_at__isnull=False,
                dogowors__deactivate_at__isnull=True,
                dogowors__balance_type="telefon"
            )

            # Исключаем текущего пользователя при редактировании
            if user_id:
                query = query.exclude(id=user_id)

            abonent_exists = query.exists()

            return Response({
                "success": True,
                "exists": abonent_exists,
                "message": "Check completed successfully"
            })
        else:
            return Response({
                "success": False, 
            }, status=400)
        
    except Exception as e:
        ic("Error in checkActiveOrNot:", str(e))
        return Response({
            "success": False,
            "error": str(e)
        }, status=500)
        
        
        

@api_view(['GET'])
def checkUniqueDogowor(request):
    ic("checkUniqueDogowor")
    try:
        # Получаем параметры из query string
        dogowor = request.GET.get('dogowor')
        dogowor_type = request.GET.get('dogowor_type')
        dogowor_id = request.GET.get('id')
        
        
        # Валидация обязательных параметров
        if not dogowor:
            return Response({
                "success": False,
                "error": "Dogowor parameter is required"
            }, status=400)
            
            
        qs = UserDogowor.objects.filter(dogowor=dogowor, balance_type=dogowor_type)
        
        if dogowor_id:
            qs = qs.exclude(id=dogowor_id)

        dogowor_exists = qs.exists()
        
        
        # dogowor_exists = UserDogowor.objects.filter(dogowor=dogowor, balance_type=dogowor_type).exists()
        # ic(dogowor_exists)
        return Response({
            "success": True,
            "exists": dogowor_exists,
            "message": "Check completed successfully"
        })
        
    except Exception as e:
        ic("Error in checkUniqueDogowor:", str(e))
        return Response({
            "success": False,
            "error": str(e)
        }, status=500)
        

@api_view(['GET'])        
def checkUniqueLogin(request):
    ic("checkUniqueLogin")
    try:
        # Получаем параметры из query string
        login = request.GET.get('login')
        login_type = request.GET.get('login_type')
        login_id = request.GET.get('id')
        
        
        # Валидация обязательных параметров
        if not login:
            return Response({
                "success": False,
                "error": "Login parameter is required"
            }, status=400)
            
            
        qs = UserDogowor.objects.filter(login=login, balance_type=login_type)
        
        if login_id:
            qs = qs.exclude(id=login_id)

        login_exists = qs.exists()
        
        
        # dogowor_exists = UserDogowor.objects.filter(dogowor=dogowor, balance_type=dogowor_type).exists()
        # ic(dogowor_exists)
        return Response({
            "success": True,
            "exists": login_exists,
            "message": "Check completed successfully"
        })
        
    except Exception as e:
        ic("Error in checkUniqueLogin:", str(e))
        return Response({
            "success": False,
            "error": str(e)
        }, status=500)
        
        

@api_view(['GET'])
def get_filtered_users(request):
    # Получаем параметры фильтрации
    searchType = request.GET.get('searchType', '')
    surname = request.GET.get('surname', '')
    name = request.GET.get('name', '')
    patronymic = request.GET.get('patronymic', '')
    is_enterprises = request.GET.get('is_enterprises', '')
    is_active = request.GET.get('is_active', '')
    etrap_id = request.GET.get('etrap', '')
    account = request.GET.get('account', '')
    hb_type = request.GET.get('hb_type', '')
    phone = request.GET.get('phone', '')
    dogowor = request.GET.get('dogowor', '')
    address = request.GET.get('address', '')
    
    # Параметры сортировки
    sort_field = request.GET.get('sort', '')
    sort_order = request.GET.get('order', 'asc')
    
    # Параметры пагинации
    page = request.GET.get('page', 1)
    page_size = request.GET.get('page_size', 10)

    try:
        page = int(page)
        page_size = int(page_size)
    except (ValueError, TypeError):
        page = 1
        page_size = 20

    qs = UserTable.objects.all()

    # Фильтруем по is_active
    if is_active == 'true':
        qs = qs.filter(
            dogowors__activate_at__isnull=False,
            dogowors__deactivate_at__isnull=True,
            dogowors__balance_type="telefon"
        ).distinct()
    elif is_active == 'false':
        qs = qs.filter(
            dogowors__deactivate_at__isnull=False,
            dogowors__balance_type="telefon"
        ).distinct()

    # Фильтруем по предприятию
    if is_enterprises:
        qs = qs.filter(is_enterprises=(is_enterprises == 'true'))

    # Фильтр по этрапу
    if etrap_id:
        qs = qs.filter(etrap__id=etrap_id)

    # Фильтры для searchType
    if searchType == 'users':
        if surname:
            qs = qs.filter(surname__icontains=surname)
        if name:
            qs = qs.filter(name__icontains=name)
        if patronymic:
            qs = qs.filter(patronymic__icontains=patronymic)
        if phone:
            qs = qs.filter(
                Q(number__icontains=phone) |
                Q(mobile_number__icontains=phone)
            )
        if dogowor:
            qs = qs.filter(dogowors__dogowor__icontains=dogowor)
    elif searchType == 'phone':
        if phone:
            qs = qs.filter(
                Q(number__icontains=phone) |
                Q(mobile_number__icontains=phone)
            )
    elif searchType == 'dogowor':
        if dogowor:
            qs = qs.filter(dogowors__dogowor__icontains=dogowor)
    elif searchType == 'address':
        if address:
            qs = qs.filter(address__icontains=address)

    # Фильтры для предприятий
    if account:
        qs = qs.filter(account=account)
    if hb_type:
        qs = qs.filter(hb_type=hb_type)

    # Применяем сортировку
    if sort_field:
        # Исправленный маппинг полей фронтенда на поля модели
        field_mapping = {
            'fullName': 'surname',  # Сортируем по фамилии для полного имени
            'number': 'number',
            'login': 'dogowors__login',
            'phone': 'mobile_number',
            'type': 'is_enterprises',
            'hbType': 'hb_type',
            'account': 'account',
            'etrap': 'etrap__etrap',  # ИСПРАВЛЕНО: используем поле 'etrap' вместо 'name'
            'address': 'address',
            'abonplata': 'abonplata',
            'id': 'id',
            'services_count': 'services_count',
            'services': 'services_count',
        }
        
        db_field = field_mapping.get(sort_field)
        if db_field:
            if sort_field in ['services', 'services_count']:
                ic("tut services_count")
                # qs = qs.annotate(services_count=Count('userservice', distinct=True))
                qs = qs.annotate(services_count=Count('userservice', filter=models.Q(userservice__is_active=True), distinct=True))
                
            if sort_order == 'desc':
                db_field = f'-{db_field}'
            
            # Для сложных полей (отношения) используем аннотацию
            if sort_field == 'login':
                # Аннотируем queryset первым логином из договоров
                qs = qs.annotate(
                    first_login=Subquery(
                        UserDogowor.objects.filter(
                            user_id=OuterRef('id'),
                            login__isnull=False
                        ).order_by('id').values('login')[:1]
                    )
                ).order_by(db_field.replace('dogowors__login', 'first_login'))
            else:
                qs = qs.order_by(db_field)
    else:
        # Сортировка по умолчанию
        qs = qs.order_by('surname', 'name')

    # Применяем distinct() и подсчитываем общее количество записей
    qs = qs.distinct()
    total_count = qs.count()

    # Создаем пагинатор
    paginator = Paginator(qs, page_size)
    
    try:
        paginated_qs = paginator.page(page)
    except EmptyPage:
        paginated_qs = paginator.page(paginator.num_pages)
        
    

    # Сериализация
    results = []
    for user in paginated_qs:
        dogowors_list = []
        for d in user.dogowors.all():
            dogowors_list.append({
                "id": d.id,
                "dogowor": d.dogowor,
                "balance": d.balance.amount if hasattr(d, "balance") else 0,
                "login": d.login,
                "balance_type": d.get_balance_type_display(),
                "balance_type2": d.balance_type,
                "comment": d.comment,
                "activate_at": d.activate_at,
                "deactivate_at": d.deactivate_at,
            })
            
        services_obj = UserService.objects.filter(user=user)
        services = []
        if services_obj.exists():
            for service in services_obj:
                if service.is_active:
                    services.append({"service": service.service.service, "price": service.actual_price})

        results.append({
            "id": user.id,
            "number": user.number,
            "surname": user.surname,
            "name": user.name,
            "patronymic": user.patronymic,
            "phone": user.mobile_number,
            "is_enterprises": user.is_enterprises,
            "account": user.account,
            "hb_type": user.hb_type,
            "etrap": str(user.etrap),
            "address": user.address,
            "abonplata": user.abonplata,
            "dogowors": dogowors_list,
            "services": services,
        })

    return JsonResponse({
        "results": results,
        "pagination": {
            "current_page": page,
            "page_size": page_size,
            "total_count": total_count,
            "total_pages": paginator.num_pages,
            "has_next": paginated_qs.has_next(),
            "has_previous": paginated_qs.has_previous(),
        },
    })



@api_view(['GET'])
def get_user_for_update_telefoniya(request):
    
    dogoworId = request.GET.get("dogoworId")
    
    dogowor_obj = get_object_or_404(UserDogowor, id=dogoworId)
    user = dogowor_obj.user

    # user_services = UserService.objects.filter(user=user)
    # for s in user_services:
    #     ic(s.connected_by.username)
    
    # ⭐⭐⭐ ПОЛУЧАЕМ ПОДРОБНЫЕ ДАННЫЕ ОБ УСЛУГАХ ⭐⭐⭐
    user_services = UserService.objects.filter(user=user, is_active=True).select_related('service')
    services_data = []
    services_ids = []
    service_dates = {}  # ← ДОБАВИТЬ ЭТО!

    for user_service in user_services:
        services_ids.append(user_service.service_id)
        services_data.append({
            "id": user_service.service_id,
            "service": user_service.service.service,
            "price": user_service.service.price,
            "actual_price": user_service.actual_price,
            "date_connected": user_service.date_connected,
            "is_active": user_service.is_active,
        })
        
        # ⭐⭐⭐ ДОБАВЛЯЕМ ДАННЫЕ В service_dates ⭐⭐⭐
        service_dates[str(user_service.service_id)] = {
            "activate_at": user_service.date_connected,  # дата подключения
            "deactivate_at": user_service.date_end if user_service.date_end else None,  # ← ИСПРАВЛЕНО: date_end вместо date_disconnected
            "comment": user_service.comment,
        }

    # ⭐⭐⭐ ФОРМИРУЕМ ДАННЫЕ ПОЛЬЗОВАТЕЛЯ ⭐⭐⭐
    user_data = {
        "id": user.id,
        "number": user.number,
        "name": user.name,
        "surname": user.surname,
        "patronymic": user.patronymic,
        "address": user.address,
        "mobile_number": user.mobile_number,
        "is_enterprises": user.is_enterprises,
        "account": user.account,
        "hb_type": user.hb_type,
        "etrap": user.etrap.id if user.etrap else None,
        "abonplata": str(user.abonplata),
        "services": services_ids,
        "services_details": services_data,
        "service_dates": service_dates,  # ← ДОБАВИТЬ ЭТО В ОТВЕТ!
        
    }
    
    # ⭐⭐⭐ ФОРМИРУЕМ ДАННЫЕ ДОГОВОРА ⭐⭐⭐
    dogowor_data = {
        "id": dogowor_obj.id,
        "dogowor": dogowor_obj.dogowor,
        "login": dogowor_obj.login,
        "balance_type": dogowor_obj.balance_type,
        "activate_at": dogowor_obj.activate_at,
        "deactivate_at": dogowor_obj.deactivate_at,
        "comment": dogowor_obj.comment,
        "already_deactivated": True if dogowor_obj.deactivate_at else False
    }
    
    # ic(dogowor_data)
    
    return Response({
        "success": True,
        "user": user_data,
        "dogowor": dogowor_data,
    })
    

@api_view(['GET'])
def get_all_services(request):
    
    services = Service.objects.filter(is_active=True)
    
    data = []
    
    for service in services:
        data.append({
            "id":service.id,
            "service":service.service,
            "price":service.price,
        })
    
    
    return Response({
        "success": True,
        "data": data,
    })
    
    
   

@api_view(['GET', 'POST'])
def export_users(request):
    export_type = request.GET.get('export_type', 'page')
    
    # Получаем параметры фильтрации из GET запроса
    searchType = request.GET.get('searchType', '')
    surname = request.GET.get('surname', '')
    name = request.GET.get('name', '')
    patronymic = request.GET.get('patronymic', '')
    is_enterprises = request.GET.get('is_enterprises', '')
    is_active = request.GET.get('is_active', '')
    etrap_id = request.GET.get('etrap', '')
    account = request.GET.get('account', '')
    hb_type = request.GET.get('hb_type', '')
    phone = request.GET.get('phone', '')
    dogowor = request.GET.get('dogowor', '')
    address = request.GET.get('address', '')
    services = request.GET.get('services', '')
    
    # ic(request.GET)

    # Базовый queryset
    qs = UserTable.objects.all()

    # Применяем фильтры (такие же как в get_filtered_users)
    if is_active == 'true':
        qs = qs.filter(
            dogowors__activate_at__isnull=False,
            dogowors__deactivate_at__isnull=True,
            dogowors__balance_type="telefon"
        ).distinct()
    elif is_active == 'false':
        qs = qs.filter(
            dogowors__deactivate_at__isnull=False,
            dogowors__balance_type="telefon"
        ).distinct()

    if is_enterprises:
        qs = qs.filter(is_enterprises=(is_enterprises == 'true'))

    if etrap_id:
        qs = qs.filter(etrap__id=etrap_id)

    if searchType == 'users':
        if surname:
            qs = qs.filter(surname__icontains=surname)
        if name:
            qs = qs.filter(name__icontains=name)
        if patronymic:
            qs = qs.filter(patronymic__icontains=patronymic)
        if phone:
            qs = qs.filter(
                Q(number__icontains=phone) |
                Q(mobile_number__icontains=phone)
            )
        if dogowor:
            qs = qs.filter(dogowors__dogowor__icontains=dogowor)
    elif searchType == 'phone':
        if phone:
            qs = qs.filter(
                Q(number__icontains=phone) |
                Q(mobile_number__icontains=phone)
            )
    elif searchType == 'dogowor':
        if dogowor:
            qs = qs.filter(dogowors__dogowor__icontains=dogowor)
    elif searchType == 'address':
        if address:
            qs = qs.filter(address__icontains=address)

    if account:
        qs = qs.filter(account=account)
    if hb_type:
        qs = qs.filter(hb_type=hb_type)

    # Обрабатываем разные типы экспорта
    if export_type == 'selected' and request.method == 'POST':
        # Для выбранных пользователей
        user_ids = request.data.get('user_ids', [])
        qs = qs.filter(id__in=user_ids)
    elif export_type == 'page':
        # Для текущей страницы - применяем пагинацию
        page = int(request.GET.get('page', 1))
        page_size = int(request.GET.get('page_size', 20))
        start = (page - 1) * page_size
        end = start + page_size
        qs = qs[start:end]
    # Для 'all' - берем все данные без пагинации

    # Создаем Excel файл
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"

    # Заголовки
    headers = [
        'ID', 'Номер', 'Фамилия', 'Имя', 'Отчество', 
        'Сотовый номер', 'Предприятие', 'Тип', 'Счет',
        'Этрап', 'Адрес', 'Абонплата', 'Договоры', 'Балансы', 'Услуги'
    ]
    
    # Стили для заголовков
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Данные
    row_num = 2
    for user in qs:
        # Получаем информацию о договорах
        dogowors_info = []
        balances_info = []
        
        for dogowor in user.dogowors.all():
            dogowors_info.append(dogowor.dogowor)
            balance = dogowor.balance.amount if hasattr(dogowor, "balance") else 0
            balances_info.append(f"{dogowor.dogowor}: {balance}")
            
        # Услуги пользователя
        services_info = []
        user_services = UserService.objects.filter(user=user, is_active=True)  # только активные услуги
        for us in user_services:
            services_info.append(f"{us.service.service} ({us.actual_price})")

        ws.cell(row=row_num, column=1, value=user.id)
        ws.cell(row=row_num, column=2, value=user.number)
        ws.cell(row=row_num, column=3, value=user.surname)
        ws.cell(row=row_num, column=4, value=user.name)
        ws.cell(row=row_num, column=5, value=user.patronymic)
        ws.cell(row=row_num, column=6, value=user.mobile_number)
        ws.cell(row=row_num, column=7, value="Да" if user.is_enterprises else "Нет")
        ws.cell(row=row_num, column=8, value=user.get_hb_type_display() if user.hb_type else "")
        ws.cell(row=row_num, column=9, value=user.account)
        ws.cell(row=row_num, column=10, value=str(user.etrap))
        ws.cell(row=row_num, column=11, value=user.address)
        ws.cell(row=row_num, column=12, value=float(user.abonplata) if user.abonplata else 0)
        ws.cell(row=row_num, column=13, value=", ".join(dogowors_info))
        ws.cell(row=row_num, column=14, value="; ".join(balances_info))
        ws.cell(row=row_num, column=15, value=", ".join(services_info))
        
        row_num += 1

    # Авто-ширина колонок
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Сохраняем в буфер
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    # Создаем HTTP ответ
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=users_export.xlsx'
    
    return response   
   
   
   
    